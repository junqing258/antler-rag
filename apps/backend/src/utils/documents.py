from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_SUFFIXES = {".txt", ".md", ".markdown", ".pdf", ".docx"}


class UnsupportedDocument(ValueError):
    pass


def extract_text(filename: str, content: bytes) -> str:
    """Extract text without invoking an external service."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise UnsupportedDocument(f"Unsupported file type '{suffix}'. Supported: {supported}")
    if suffix in {".txt", ".md", ".markdown"}:
        return content.decode("utf-8", errors="replace").strip()
    if suffix == ".pdf":
        reader = PdfReader(BytesIO(content))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()
    document = DocxDocument(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = [
        "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        for table in document.tables
    ]
    return "\n\n".join(paragraphs + tables).strip()
